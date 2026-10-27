"""Exploratory five-class baseline on the published Peanut-seed spreadsheet.

This is a separate spectral experiment, not a test of the mobile image model.
The source provides no group IDs or explicit numeric-label mapping.
"""
import argparse
import hashlib
import json
import platform
from collections import Counter
from pathlib import Path

import lightgbm as lgb
import numpy as np
import sklearn
from openpyxl import load_workbook
from sklearn.dummy import DummyClassifier
from sklearn.metrics import accuracy_score, confusion_matrix, f1_score, classification_report
from sklearn.model_selection import StratifiedGroupKFold

EXPECTED_SHA256 = '5ae0b3ae5175e5adebfb5091045cacb02f05c9e50593ac0a7cf4fadd06f28866'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('xlsx', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    digest = hashlib.sha256(args.xlsx.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        raise ValueError('Unexpected dataset hash; inspect the source before proceeding')
    book = load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        if book.sheetnames != ['Sheet1']:
            raise ValueError('Unexpected workbook structure')
        matrix = np.asarray(list(book['Sheet1'].values), dtype=np.float64)
    finally:
        book.close()
    if matrix.shape != (2200, 237) or not np.isfinite(matrix).all():
        raise ValueError(f'Unexpected shape or non-finite values: {matrix.shape}')
    x = matrix[:, :-1]
    y = matrix[:, -1].astype(int)
    if not np.array_equal(y, matrix[:, -1]) or set(y) != {1, 2, 3, 4, 5}:
        raise ValueError('Unexpected numeric labels')
    # Group exact duplicate spectra to prevent identical rows crossing splits.
    _, groups, counts = np.unique(x, axis=0, return_inverse=True, return_counts=True)
    duplicate_count = int(np.sum(counts - 1))
    if any(np.unique(y[groups == group]).size > 1 for group in np.flatnonzero(counts > 1)):
        raise ValueError('Identical spectra have conflicting labels')
    splitter = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=42)
    train, test = next(splitter.split(x, y, groups=groups))
    assert not set(groups[train]).intersection(groups[test])
    models = {
        'majority_class': DummyClassifier(strategy='most_frequent'),
        'lightgbm': lgb.LGBMClassifier(n_estimators=100, learning_rate=0.05,
            num_leaves=15, random_state=42, n_jobs=2, verbosity=-1),
    }
    results = {}
    for name, model in models.items():
        model.fit(x[train], y[train])
        pred = model.predict(x[test])
        results[name] = {
            'accuracy': accuracy_score(y[test], pred),
            'macro_f1': f1_score(y[test], pred, average='macro', zero_division=0),
            'confusion_matrix': confusion_matrix(y[test], pred, labels=[1, 2, 3, 4, 5]).tolist(),
            'classification_report': classification_report(y[test], pred, labels=[1, 2, 3, 4, 5], output_dict=True, zero_division=0),
        }
    report = {
        'dataset_url': 'https://github.com/wuqingsongwj/Peanut-seed',
        'paper_url': 'https://doi.org/10.3389/fpls.2022.1047479',
        'versions': {'python': platform.python_version(), 'numpy': np.__version__,
                     'scikit_learn': sklearn.__version__, 'lightgbm': lgb.__version__},
        'sha256': digest, 'rows': len(y), 'numeric_input_columns': x.shape[1],
        'class_counts': dict(sorted(Counter(int(v) for v in y).items())),
        'exact_duplicate_feature_rows': duplicate_count,
        'split': {'type': 'stratified exact-spectrum-group five-fold, first fold held out', 'seed': 42, 'train': len(train), 'test': len(test)},
        'limitations': ['Numeric labels are not mapped to class names in the workbook.',
            'No sample, acquisition-session or batch identifiers; group leakage cannot be excluded.',
            '236 input columns observed; paper describes 235 spectral bands.',
            'Five-class spectral model is separate from the 16-feature mobile image model.',
            'Not a validation of aflatoxin concentration or food safety.'],
        'results': results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({k: {'accuracy': v['accuracy'], 'macro_f1': v['macro_f1']} for k, v in results.items()}, indent=2))


if __name__ == '__main__':
    main()
